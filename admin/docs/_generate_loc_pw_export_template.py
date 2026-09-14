# -*- coding: utf-8 -*-
"""
生成「本地 / 外州私仓发货」列表导出模板。

用法：
  python admin/docs/_generate_loc_pw_export_template.py

输出：
  admin/templates/本地私仓发货导出模板.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent.parent / "templates"
OUT = OUT_DIR / "本地私仓发货导出模板.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
KEY_FILL = PatternFill("solid", fgColor="C2410C")
SAMPLE_MERGE_FILL = PatternFill("solid", fgColor="FFF7ED")
SAMPLE_NORMAL_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SECTION_FONT = Font(bold=True, size=11, color="1F4E79")
HINT_FONT = Font(size=10, color="666666")
BODY_FONT = Font(size=11, color="1E293B")
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 与列表字段对齐，并新增「发车类型」「合并发货BOL号」
# 关键列：发车类型、BOL号、合并发货BOL号、里程
COLUMNS = [
    ("发货模式", "列表已有", "普通发货 / 拆分发货 / 合并发货。合并子票按「合并发货」导出，不导出父单汇总行。"),
    ("发车类型", "新增", "FTL发车 / LTL发车。取该票已填写的发车类型；合并票与父单同一车次时沿用父单。未安排出库可留空。"),
    ("BOL号", "列表已有", "独立票号。普通/拆分导出本票 BOL；合并导出各子票原 BOL。"),
    ("合并发货BOL号", "新增", "仅合并发货填写父单 BOL（如 BOLO2607099001）。普通/拆分留空。同一合并组各子票此列相同。"),
    ("Customer Ref No", "列表已有", "该票客户参考号。"),
    ("状态", "列表已有", "待处理 / 处理中 / 待取货 / 运输中 / 已签收 / 暂缓处理 / 退仓待执行 等。"),
    ("客户", "列表已有", "该票客户名称。"),
    ("柜号", "列表已有", "该票柜号。"),
    ("到仓日期", "列表已有", "yyyy-mm-dd"),
    ("拆柜日期", "列表已有", "yyyy-mm-dd"),
    ("库位", "列表已有", "该票库位。"),
    ("Address", "列表已有", "该票送仓地址。"),
    ("QTY(CTNS)", "列表已有", "该票件数。"),
    ("Vol(CBM)", "列表已有", "该票体积。"),
    ("GW(LBS)", "列表已有", "该票毛重。"),
    ("城市", "列表已有", ""),
    ("州", "列表已有", ""),
    ("邮编", "列表已有", ""),
    ("里程(miles)", "列表已有", "该票独立里程。合并时取各子票距离，不用父单「—」。"),
    ("公司名称", "列表已有", ""),
    ("联系人", "列表已有", ""),
    ("手机", "列表已有", ""),
    ("邮箱", "列表已有", ""),
    ("预估板数", "列表已有", ""),
    ("实际板数", "列表已有", ""),
    ("最早送仓日期", "列表已有", "yyyy-mm-dd"),
    ("最晚送仓日期", "列表已有", "yyyy-mm-dd"),
    ("中文品名", "列表已有", ""),
    ("英文品名", "列表已有", ""),
    ("拆柜要求", "列表已有", ""),
    ("预约要求", "列表已有", ""),
    ("预约文件", "列表已有", "文件名，多个用空格或分号分隔。"),
    ("私卡成本价", "列表已有", ""),
    ("内部备注", "列表已有", "系统内部备注。"),
    ("暂缓处理原因", "列表已有", ""),
    ("发车时间", "列表已有", "yyyy-mm-dd HH:mm"),
    ("发车凭证", "列表已有", ""),
    ("预计送达时间", "列表已有", "yyyy-mm-dd 或 yyyy-mm-dd HH:mm"),
    ("签收时间", "列表已有", "yyyy-mm-dd HH:mm"),
    ("POD", "列表已有", ""),
    ("货件ID", "列表已有", "该票独立货件 ID。"),
    ("系统单号", "列表已有", "该票独立系统单号。"),
    ("创建人", "列表已有", ""),
    ("创建时间", "列表已有", "yyyy-mm-dd HH:mm:ss"),
]

KEY_COLS = {"发车类型", "BOL号", "合并发货BOL号", "里程(miles)"}
COL_NAMES = [name for name, _, _ in COLUMNS]

# 示例：一行一票。合并组拆成 3 行，各自带单号、里程与合并发货BOL号。
SAMPLE_ROWS = [
    {
        "发货模式": "普通发货",
        "发车类型": "",
        "BOL号": "BOLO2607090412",
        "Customer Ref No": "ref-pending-01",
        "状态": "待处理",
        "客户": "Nova Parts Inc.",
        "柜号": "MSKU8811001",
        "到仓日期": "2026-04-30",
        "拆柜日期": "2026-05-01",
        "库位": "F-02-01",
        "Address": "5100 Etiwanda Ave",
        "QTY(CTNS)": 20,
        "Vol(CBM)": 0.95,
        "GW(LBS)": 140,
        "城市": "Jurupa Valley",
        "州": "CA",
        "邮编": "91752",
        "里程(miles)": 58,
        "联系人": "Nick",
        "手机": "951-555-2210",
        "邮箱": "nick@novaparts.example.com",
        "预估板数": 2,
        "实际板数": 2,
        "最早送仓日期": "2026-05-02",
        "最晚送仓日期": "2026-05-05",
        "私卡成本价": 68,
        "货件ID": "TLP2606230412-0001",
        "系统单号": "TLP2606230412",
        "创建人": "张伟",
        "创建时间": "2026-04-30 09:10:00",
    },
    {
        "发货模式": "合并发货",
        "发车类型": "",
        "BOL号": "BOLO2607090401",
        "合并发货BOL号": "BOLO2607099001",
        "Customer Ref No": "ref-001-customerX, ref-001b-customerX",
        "状态": "待处理",
        "客户": "ABC Trading Co.",
        "柜号": "MSKU1234567",
        "到仓日期": "2026-04-26",
        "拆柜日期": "2026-04-27",
        "库位": "A-01-01",
        "Address": "1234 Warehouse Blvd",
        "QTY(CTNS)": 28,
        "Vol(CBM)": 1.10,
        "GW(LBS)": 210,
        "城市": "Ontario",
        "州": "CA",
        "邮编": "91761",
        "里程(miles)": 42,
        "联系人": "Tom",
        "手机": "626-111-0001",
        "邮箱": "tom.x@example.com",
        "预估板数": 2,
        "实际板数": 2,
        "最早送仓日期": "2026-04-27",
        "最晚送仓日期": "2026-04-29",
        "预约要求": "需提前预约卸货口",
        "预约文件": "appt-0401-A.pdf dock-req-0401.xlsx appt-0401-A2.pdf",
        "私卡成本价": 40,
        "预计送达时间": "2026-04-28",
        "货件ID": "TLP2606230401-0001 TLP2606230401-0002",
        "系统单号": "TLP2606230401",
        "创建人": "张伟",
        "创建时间": "2026-04-25 10:00:00",
    },
    {
        "发货模式": "合并发货",
        "发车类型": "",
        "BOL号": "BOLO2607090391",
        "合并发货BOL号": "BOLO2607099001",
        "Customer Ref No": "ref-002-customerY",
        "状态": "待处理",
        "客户": "Beta Logistics Inc.",
        "柜号": "MSKU2233445",
        "到仓日期": "2026-04-26",
        "拆柜日期": "2026-04-27",
        "库位": "A-01-02",
        "Address": "892 Carrier Row",
        "QTY(CTNS)": 25,
        "Vol(CBM)": 0.90,
        "GW(LBS)": 125,
        "城市": "Long Beach",
        "州": "CA",
        "邮编": "90802",
        "里程(miles)": 35,
        "联系人": "Jane",
        "手机": "562-222-0002",
        "邮箱": "jane.y@example.com",
        "预估板数": 2,
        "实际板数": 2,
        "最早送仓日期": "2026-04-27",
        "最晚送仓日期": "2026-04-30",
        "私卡成本价": 45,
        "预计送达时间": "2026-04-28",
        "货件ID": "TLP2606230391-0001",
        "系统单号": "TLP2606230391",
        "创建人": "李晓华",
        "创建时间": "2026-04-25 10:05:00",
    },
    {
        "发货模式": "合并发货",
        "发车类型": "",
        "BOL号": "BOLO2607090392-1",
        "合并发货BOL号": "BOLO2607099001",
        "Customer Ref No": "ref-003-customerZ",
        "状态": "待处理",
        "客户": "Gamma Retail LLC",
        "柜号": "MSKU3390008",
        "到仓日期": "2026-04-26",
        "拆柜日期": "2026-04-27",
        "库位": "A-02-06",
        "Address": "4560 Milliken Ave",
        "QTY(CTNS)": 20,
        "Vol(CBM)": 1.35,
        "GW(LBS)": 125,
        "城市": "Ontario",
        "州": "CA",
        "邮编": "91761",
        "里程(miles)": 42,
        "联系人": "Mike",
        "手机": "626-333-0003",
        "邮箱": "mike.z@example.com",
        "预估板数": 1,
        "实际板数": 1,
        "最早送仓日期": "2026-04-28",
        "最晚送仓日期": "2026-04-29",
        "私卡成本价": 35,
        "预计送达时间": "2026-04-29",
        "货件ID": "TLP2606230392-0001",
        "系统单号": "TLP2606230392",
        "创建人": "张伟",
        "创建时间": "2026-04-25 10:10:00",
    },
    {
        "发货模式": "普通发货",
        "发车类型": "LTL发车",
        "BOL号": "BOLO2607090403",
        "Customer Ref No": "ref-009ff",
        "状态": "待取货",
        "客户": "Echo Supply Co.",
        "柜号": "MSKU3390001",
        "到仓日期": "2026-04-27",
        "拆柜日期": "2026-04-28",
        "库位": "B-02-01",
        "Address": "5678 Commerce Way",
        "QTY(CTNS)": 35,
        "Vol(CBM)": 1.55,
        "GW(LBS)": 210,
        "城市": "Rancho Cucamonga",
        "州": "CA",
        "邮编": "91730",
        "里程(miles)": 38,
        "联系人": "Mike",
        "手机": "909-000-3300",
        "邮箱": "mike@example.com",
        "预估板数": 3,
        "实际板数": 3,
        "最早送仓日期": "2026-04-28",
        "最晚送仓日期": "2026-04-30",
        "私卡成本价": 150,
        "预计送达时间": "2026-04-29",
        "货件ID": "TLP2606230403-0001",
        "系统单号": "TLP2606230403",
        "创建人": "王芳",
        "创建时间": "2026-04-26 09:20:00",
    },
    {
        "发货模式": "拆分发货",
        "发车类型": "LTL发车",
        "BOL号": "BOLO2607090402-2",
        "Customer Ref No": "HK-2026-0402",
        "状态": "运输中",
        "客户": "Fox Brands Ltd.",
        "柜号": "MSKU2234567",
        "到仓日期": "2026-04-26",
        "拆柜日期": "2026-04-27",
        "库位": "A-01-03",
        "Address": "1234 Warehouse Blvd",
        "QTY(CTNS)": 14,
        "Vol(CBM)": 1.10,
        "GW(LBS)": 168,
        "城市": "Ontario",
        "州": "CA",
        "邮编": "91761",
        "里程(miles)": 42,
        "联系人": "Lucy",
        "手机": "626-000-0001",
        "邮箱": "lucy@example.com",
        "预估板数": 2,
        "实际板数": 2,
        "最早送仓日期": "2026-04-27",
        "最晚送仓日期": "2026-04-29",
        "私卡成本价": 105,
        "发车时间": "2026-04-28 08:00",
        "预计送达时间": "2026-04-28",
        "货件ID": "TLP2606230402-0001",
        "系统单号": "TLP2606230402",
        "创建人": "李晓华",
        "创建时间": "2026-04-25 10:15:00",
    },
    {
        "发货模式": "普通发货",
        "发车类型": "FTL发车",
        "BOL号": "BOLO2607090408",
        "Customer Ref No": "ref-sig01",
        "状态": "已签收",
        "客户": "Gamma Retail LLC",
        "柜号": "MSKU8899001",
        "到仓日期": "2026-04-28",
        "拆柜日期": "2026-04-29",
        "库位": "C-03-01",
        "Address": "9100 Industrial Pkwy",
        "QTY(CTNS)": 18,
        "Vol(CBM)": 0.95,
        "GW(LBS)": 142,
        "城市": "Fontana",
        "州": "CA",
        "邮编": "92335",
        "里程(miles)": 55,
        "联系人": "Amy",
        "手机": "909-100-7788",
        "邮箱": "amy@example.com",
        "预估板数": 1,
        "实际板数": 1,
        "最早送仓日期": "2026-04-29",
        "最晚送仓日期": "2026-05-02",
        "私卡成本价": 88,
        "发车时间": "2026-04-29 09:15",
        "预计送达时间": "2026-05-01",
        "签收时间": "2026-04-30 14:20",
        "POD": "已上传",
        "货件ID": "TLP2606230408-0001",
        "系统单号": "TLP2606230408",
        "创建人": "系统",
        "创建时间": "2026-04-28 16:00:00",
    },
]


def style_header_row(ws, row, col_count, key_indexes=None):
    key_indexes = key_indexes or set()
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = KEY_FILL if c in key_indexes else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize_columns(ws, min_width=10, max_width=42):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col_idx).value
            if v is None:
                continue
            first = str(v).split("\n")[0]
            max_len = max(max_len, len(first))
        ws.column_dimensions[letter].width = min(max(max_len + 2, min_width), max_width)


def sheet_instructions(wb):
    ws = wb.active
    ws.title = "导出说明"
    ws["A1"] = "本地 / 外州私仓发货 · 列表导出模板"
    ws["A1"].font = TITLE_FONT
    lines = [
        "",
        "【业务规则】",
        "1. 一行一票：普通发货、拆分发货、合并发货都按独立 BOL 各导出一行。",
        "2. 单号与距离独立：每行写该票自己的 BOL号、系统单号、货件ID、里程(miles)。",
        "3. 合并发货同时给出父单号：各子票一行，「BOL号」= 子票原单，「合并发货BOL号」= 父单（如 BOLO2607099001）。",
        "4. 不导出父单汇总行（父单地址/里程常为「—」）。三个子票就出三行，各带自己的单号和里程。",
        "5. 发车类型：FTL发车 / LTL发车。已安排出库或之后的票填写；合并票与父单同一车次时沿用父单。",
        "6. 合板 ≠ 合并发货：仅打板合并仍按普通/拆分独立票导出，「合并发货BOL号」留空。",
        "",
        "【橙色表头】发车类型、BOL号、合并发货BOL号、里程(miles) —— 本次业务重点列。",
        "",
        "【Sheet】",
        "· 字段说明：列名、是否列表已有、填写说明。",
        "· 示例数据：可直接对照列表演示数据；「下载列表信息」按本表列序导出。",
    ]
    for i, text in enumerate(lines, start=2):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = SECTION_FONT if text.startswith("【") else (HINT_FONT if "橙色" in text or text.startswith("·") else BODY_FONT)
        cell.alignment = WRAP
    ws.column_dimensions["A"].width = 118
    ws.row_dimensions[1].height = 24
    for r in range(3, 9):
        ws.row_dimensions[r].height = 20


def sheet_fields(wb):
    ws = wb.create_sheet("字段说明")
    ws["A1"] = "导出字段"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:C1")
    headers = ["字段名", "来源", "说明"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=h)
    style_header_row(ws, 3, 3)
    for i, (name, source, hint) in enumerate(COLUMNS, start=4):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=source)
        ws.cell(row=i, column=3, value=hint)
        for c in range(1, 4):
            cell = ws.cell(row=i, column=c)
            cell.alignment = WRAP
            cell.border = BORDER
            cell.font = BODY_FONT
            if name in KEY_COLS:
                cell.fill = PatternFill("solid", fgColor="FFEDD5")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 88
    ws.freeze_panes = "A4"


def sheet_sample(wb):
    ws = wb.create_sheet("示例数据")
    key_indexes = {i + 1 for i, name in enumerate(COL_NAMES) if name in KEY_COLS}
    for c, name in enumerate(COL_NAMES, start=1):
        ws.cell(row=1, column=c, value=name)
    style_header_row(ws, 1, len(COL_NAMES), key_indexes)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COL_NAMES))}{len(SAMPLE_ROWS) + 1}"

    for r_idx, row in enumerate(SAMPLE_ROWS, start=2):
        fill = SAMPLE_MERGE_FILL if row.get("发货模式") == "合并发货" else SAMPLE_NORMAL_FILL
        for c, name in enumerate(COL_NAMES, start=1):
            cell = ws.cell(row=r_idx, column=c, value=row.get(name, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.fill = fill
        ws.row_dimensions[r_idx].height = 36

    widths = {
        "发货模式": 12,
        "发车类型": 12,
        "BOL号": 20,
        "合并发货BOL号": 20,
        "Customer Ref No": 36,
        "状态": 10,
        "客户": 22,
        "柜号": 14,
        "Address": 24,
        "邮箱": 28,
        "预约文件": 28,
        "货件ID": 28,
        "系统单号": 16,
        "创建时间": 20,
    }
    for i, name in enumerate(COL_NAMES, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheet_instructions(wb)
    sheet_fields(wb)
    sheet_sample(wb)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
