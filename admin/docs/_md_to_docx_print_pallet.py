# -*- coding: utf-8 -*-
"""One-off: convert 提拆派-打印板标-需求说明.md -> .docx"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "提拆派-打印板标-需求说明.md"
DST = ROOT / "提拆派-打印板标-需求说明.docx"


def set_run_font(run, size=None, bold=None, code=False):
    run.font.name = "Consolas" if code else "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, content):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=10, code=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def shade_header_cell(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F5F9")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_inline(p, val)
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
            if i == 0:
                shade_header_cell(cell)


def parse_table(block_lines):
    rows = []
    for line in block_lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def main():
    text = SRC.read_text(encoding="utf-8")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for sname in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            hs = doc.styles[sname]
            hs.font.name = "微软雅黑"
            hs._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            hs.font.color.rgb = RGBColor(0x1A, 0x20, 0x35)
        except Exception:
            pass
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, parse_table(table_buf))
            table_buf = []

    def flush_code():
        nonlocal code_buf, in_code
        if code_buf:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.3)
            run = p.add_run("\n".join(code_buf))
            set_run_font(run, size=9, code=True)
            code_buf = []
        in_code = False

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
            else:
                flush_table()
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_buf.append(lines[i])
                i += 1
            flush_table()
            continue

        flush_table()

        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            for run in p.runs:
                set_run_font(run, size=18, bold=True)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("---"):
            pass
        elif re.match(r"^[-*]\s+", line.strip()) or re.match(r"^\d+\.\s+", line.strip()):
            content = re.sub(r"^[-*]\s+", "", line.strip())
            content = re.sub(r"^\d+\.\s+", "", content)
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, content)
        else:
            p = doc.add_paragraph()
            add_inline(p, line.strip())
        i += 1

    flush_table()
    flush_code()
    doc.save(DST)
    print("OK:", DST)
    print("size:", DST.stat().st_size)


if __name__ == "__main__":
    main()
