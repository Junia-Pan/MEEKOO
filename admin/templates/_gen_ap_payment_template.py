# -*- coding: utf-8 -*-
"""Generate 付款凭证导入模板.xlsx — run once, then delete or keep for regen."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

out = Path(__file__).with_name("付款凭证导入模板.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "付款凭证"

# 第 1 行：各字段填写提示（导入解析时跳过）
# 第 2 行：表头
# 第 3 行起：数据（含示例）
tips = [
    "【必填】填银行回单/支付平台上的流水号；同一「付款银行账号」下不能重复",
    "【必填】填系统里已有的供应商全称，须一字不差",
    "【必填】只能填：USD、CNY、HKD（三选一）",
    "【必填】填正数金额，最多两位小数，如 220.00",
    "【必填】填到秒，格式：yyyy-mm-dd HH:mm:ss\n例：2026-10-28 11:00:00",
    "【必填】填系统「银行账号」里的展示名称；与流水号一起判重",
    "【选填】可空。凭证号系统生成；水单请导入后在系统里补传",
]
headers = [
    "交易流水号*",
    "供应商*",
    "币种*",
    "付款金额*",
    "付款时间*",
    "付款银行账号*",
    "备注",
]
samples = [
    ["TXN20261028001", "速达运输", "USD", 220.00, "2026-10-28 11:00:00", "BANK OF AMERICA + ***9180", "10月卡派"],
    ["TXN20261028002", "平安仓储", "USD", 50.50, "2026-10-28 14:30:00", "BANK OF AMERICA + ***9180", ""],
    ["TXN20261028003", "快马物流", "CNY", 1280.00, "2026-10-27 09:15:30", "Chase-****1234", "示例备注，导入前请删除示例行"],
]

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True, size=11)
tip_fill = PatternFill("solid", fgColor="FFF7ED")
tip_font = Font(color="9A3412", size=10)
req_fill = PatternFill("solid", fgColor="FFF2CC")
sample_fill = PatternFill("solid", fgColor="F3F4F6")
thin = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)

for col, tip in enumerate(tips, 1):
    cell = ws.cell(1, col, tip)
    cell.fill = tip_fill
    cell.font = tip_font
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = thin

for col, h in enumerate(headers, 1):
    cell = ws.cell(2, col, h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin

for r, row in enumerate(samples, 3):
    for c, val in enumerate(row, 1):
        cell = ws.cell(r, c, val)
        cell.border = thin
        cell.fill = sample_fill
        if c == 4:
            cell.number_format = "0.00"
        cell.alignment = Alignment(vertical="center")

for r in range(6, 56):
    for c in range(1, 8):
        cell = ws.cell(r, c, "")
        cell.border = thin
        if c <= 6:
            cell.fill = req_fill

widths = [28, 30, 14, 16, 22, 32, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 48
ws.row_dimensions[2].height = 28
ws.freeze_panes = "A3"
ws.auto_filter.ref = "A2:G2"

dv_currency = DataValidation(type="list", formula1='"USD,CNY,HKD"', allow_blank=True)
dv_currency.error = "请填写 USD / CNY / HKD"
dv_currency.errorTitle = "币种无效"
ws.add_data_validation(dv_currency)
dv_currency.add("C3:C500")

ws3 = wb.create_sheet("参考-可选值")
ws3["A1"] = "币种"
ws3["B1"] = "付款银行账号示例（以系统配置为准）"
ws3["C1"] = "供应商示例（以系统配置为准）"
for col in range(1, 4):
    ws3.cell(1, col).fill = header_fill
    ws3.cell(1, col).font = header_font
    ws3.cell(1, col).border = thin

for i, ccy in enumerate(["USD", "CNY", "HKD"], 2):
    ws3.cell(i, 1, ccy).border = thin

banks = ["BANK OF AMERICA + ***9180", "BOA-****5678", "Chase-****1234"]
suppliers = ["速达运输", "平安仓储", "快马物流"]
for i, (b, s) in enumerate(zip(banks, suppliers), 2):
    ws3.cell(i, 2, b).border = thin
    ws3.cell(i, 3, s).border = thin

ws3.column_dimensions["A"].width = 10
ws3.column_dimensions["B"].width = 36
ws3.column_dimensions["C"].width = 18

wb.save(out)
print("saved", out, "sheets", wb.sheetnames, "size", out.stat().st_size)
