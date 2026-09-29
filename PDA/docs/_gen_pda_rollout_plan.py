# -*- coding: utf-8 -*-
"""Generate PDA feature rollout plan Excel (2 sheets: summary + detail)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
import os

wb = Workbook()

thin = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
title_font = Font(name="微软雅黑", bold=True, size=16, color="1F4E79")
section_font = Font(name="微软雅黑", bold=True, size=12, color="1F4E79")
normal_font = Font(name="微软雅黑", size=10)
muted_font = Font(name="微软雅黑", size=10, color="666666")
wrap = Alignment(wrap_text=True, vertical="center", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")
phase_fills = {
    "P0": PatternFill("solid", fgColor="C6EFCE"),
    "P1": PatternFill("solid", fgColor="FFF2CC"),
    "P2": PatternFill("solid", fgColor="FCE4D6"),
    "P3": PatternFill("solid", fgColor="DDEBF7"),
    "P4": PatternFill("solid", fgColor="E2D5F1"),
}
status_fills = {
    "已上线使用": PatternFill("solid", fgColor="C6EFCE"),
    "待推广": PatternFill("solid", fgColor="FFF2CC"),
    "已完成": PatternFill("solid", fgColor="C6EFCE"),
    "未开始": PatternFill("solid", fgColor="F2F2F2"),
}


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin


def style_body(ws, start_row, end_row, cols):
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = normal_font
            cell.alignment = wrap
            cell.border = thin


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_merged(ws, row, cols, text, font=None, height=None):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    ws.cell(row=row, column=1).font = font or normal_font
    ws.cell(row=row, column=1).alignment = wrap
    if height:
        ws.row_dimensions[row].height = height


# ==================== Sheet 1: 汇总 ====================
ws0 = wb.active
ws0.title = "汇总"

write_merged(ws0, 1, 6, "Meekoo WMS · PDA 功能推广使用计划（汇总）", title_font, 28)
write_merged(
    ws0,
    2,
    6,
    f"编制日期：{date.today().isoformat()}　｜　现状：仅「自提出库」已在用；下一优先：「到仓确认」",
    muted_font,
)

write_merged(ws0, 4, 6, "一、推广目标", section_font)
write_merged(
    ws0,
    5,
    6,
    "以自提出库为样板，优先打通入库（到仓确认 → 拆柜打板 → 入库上架），再推常规出库与异常逆向；"
    "一线以 PDA 扫码为主、PC 为辅，提升到仓时效与账实可追溯。",
    normal_font,
    40,
)

write_merged(ws0, 7, 6, "二、分阶段节奏", section_font)
headers0 = ["阶段", "周期", "推广模块", "核心目标", "成功标准", "前置依赖"]
for i, h in enumerate(headers0, 1):
    ws0.cell(row=8, column=i, value=h)
style_header(ws0, 8, 6)

phases = [
    ("P0 已落地", "已完成", "自提出库", "上门提货扫板放行，形成使用习惯", "自提 100% PDA 放行", "账号/设备已就绪"),
    ("P1 入库切入", "第 1–2 周", "到仓确认", "现场扫码+实收板数+拍照，替代补录", "试点 PDA 占比≥80%；板数不符必登记", "单号可扫；照片可传"),
    ("P2 入库闭环", "第 3–5 周", "拆柜打板 → 入库上架", "板标打印+库位绑定可追溯", "新板 100% 有板标+库位", "打印机；库位码；板标规则"),
    ("P3 出库扩展", "第 6–8 周", "出库备货 → 确认发车；快递出库", "卡车/快递出库扫码放行", "错拣下降；发车扣库存及时", "备货/发车单可查；通知联调"),
    ("P4 异常逆向", "第 9–10 周", "问题件/工单/退仓/板标重打", "异常当场登记、逆向可闭环", "问题件当日登记≥90%", "工单联动；退仓字典"),
]
for i, row in enumerate(phases):
    r = 9 + i
    for c, v in enumerate(row, 1):
        ws0.cell(row=r, column=c, value=v)
    style_body(ws0, r, r, 6)
    key = row[0][:2]
    if key in phase_fills:
        ws0.cell(row=r, column=1).fill = phase_fills[key]
    ws0.row_dimensions[r].height = 36

write_merged(ws0, 15, 6, "三、推广原则", section_font)
for i, p in enumerate([
    "1. 先高频后低频：先到仓确认，再拆柜/上架，最后异常逆向。",
    "2. 先试点后全仓：选 1 仓试点 1–2 周，再复制。",
    "3. 双轨有期限：过渡期可 PC 兜底，设强制切换日，避免长期双轨。",
    "4. 复用自提经验：账号、扫码、拍照、二次确认沿用现有 SOP。",
    "5. 培训+跟班：新模块首周关键人跟班纠错。",
]):
    write_merged(ws0, 16 + i, 6, p)

write_merged(ws0, 22, 6, "四、组织分工", section_font)
for i, h in enumerate(["角色", "职责", "产出物"], 1):
    ws0.cell(row=23, column=i, value=h)
style_header(ws0, 23, 3)
for i, row in enumerate([
    ("产品/业务负责人", "定范围、节奏、成功标准", "推广计划、切换公告"),
    ("仓主管（试点仓）", "种子用户、监督执行", "每日使用率简报"),
    ("IT/实施", "账号、设备、打印机、网络", "设备清单、开通记录"),
    ("培训讲师（可兼主管）", "培训、SOP、跟班", "签到、考核记录"),
    ("一线操作员", "按 SOP 使用 PDA", "现场作业数据"),
]):
    r = 24 + i
    for c, v in enumerate(row, 1):
        ws0.cell(row=r, column=c, value=v)
    style_body(ws0, r, r, 3)
    ws0.row_dimensions[r].height = 24

write_merged(ws0, 30, 6, "五、关键里程碑", section_font)
for i, h in enumerate(["里程碑", "计划周", "完成标准", "状态", "实际日期", "备注"], 1):
    ws0.cell(row=31, column=i, value=h)
style_header(ws0, 31, 6)
miles = [
    ("M0 自提出库稳定运行", "已完成", "持续使用、无重大投诉", "已完成", "", ""),
    ("M1 试点仓到仓培训完成", "W1", "到仓岗 100% 实操达标", "未开始", "", ""),
    ("M2 到仓 PDA 占比≥80%", "W2", "连续 3 个工作日达标", "未开始", "", ""),
    ("M3 到仓强制切换", "W2 末", "PC 补录需主管审批", "未开始", "", ""),
    ("M4 入库闭环试点完成", "W5", "新板有板标+库位", "未开始", "", ""),
    ("M5 常规出库试点完成", "W8", "备货+发车主路径 PDA 化", "未开始", "", ""),
    ("M6 异常逆向演练完成", "W10", "问题件/退仓案例通过", "未开始", "", ""),
    ("M7 多仓复制完成", "视仓数", "各仓检查清单与切换完成", "未开始", "", ""),
]
for i, row in enumerate(miles):
    r = 32 + i
    for c, v in enumerate(row, 1):
        ws0.cell(row=r, column=c, value=v)
    style_body(ws0, r, r, 6)
    if row[3] in status_fills:
        ws0.cell(row=r, column=4).fill = status_fills[row[3]]
    ws0.row_dimensions[r].height = 24

write_merged(ws0, 41, 6, "六、主要风险（摘要）", section_font)
for i, h in enumerate(["风险", "应对", "责任人"], 1):
    ws0.cell(row=42, column=i, value=h)
style_header(ws0, 42, 3)
for i, row in enumerate([
    ("仍回 PC 补录", "强制切换日 + PC 补录需审批 + 种子用户带教", "仓主管"),
    ("条码扫不上", "上线前抽测；手输校验；补贴清晰标签", "IT"),
    ("高峰机不够", "按并发配机；错峰充电；自提机复用", "IT+主管"),
    ("板数不符不登记", "系统强提示 + 抽检纳入考核", "产品+主管"),
    ("双轨过久两套账", "明确切换日；限制/关闭 PC 入口", "产品"),
]):
    r = 43 + i
    for c, v in enumerate(row, 1):
        ws0.cell(row=r, column=c, value=v)
    style_body(ws0, r, r, 3)
    ws0.row_dimensions[r].height = 24

set_widths(ws0, [28, 12, 36, 36, 14, 18])
ws0.freeze_panes = "A9"

# ==================== Sheet 2: 明细 ====================
ws1 = wb.create_sheet("明细")

write_merged(ws1, 1, 10, "Meekoo WMS · PDA 功能推广使用计划（明细）", title_font, 28)
write_merged(
    ws1,
    2,
    10,
    "模块清单按作业操作顺序排列（入库→出库→异常→辅助）；另含到仓确认执行明细、后续阶段动作与上线检查项。",
    muted_font,
)

# --- A. 模块清单（按作业操作顺序：入库 → 出库 → 异常 → 辅助）---
write_merged(ws1, 4, 10, "A. 模块清单（按作业操作顺序）", section_font)
mod_h = [
    "操作顺序", "业务域", "模块", "页面", "功能摘要", "当前状态",
    "阶段", "优先级", "岗位", "备注",
]
for i, h in enumerate(mod_h, 1):
    ws1.cell(row=5, column=i, value=h)
style_header(ws1, 5, 10)

# 顺序：到仓→拆柜打板→测量→上架→备货→发车→自提→快递→异常→辅助
modules = [
    (1, "入库", "到仓确认", "inbound.html", "扫柜号/单号→实收板数→拍照确认；不符走问题件", "待推广", "P1", "高（下一优先）", "到仓岗", "入库入口，建议首推"),
    (2, "入库", "拆柜打板", "print-pallet.html", "生成板标，蓝牙/Wi-Fi 打印", "待推广", "P2", "高", "拆柜员", "依赖打印机与板标规则"),
    (3, "入库", "私卡测量", "measurement.html", "上架前测尺寸重量（外州私仓）", "待推广", "P2", "中（按仓）", "测量员", "非全仓必推；穿插在上架前"),
    (4, "入库", "入库上架", "putaway.html", "扫板标+库位码绑定", "待推广", "P2", "高", "上架员", "与打板衔接闭环"),
    (5, "出库", "出库备货", "outbound-pick.html", "扫备货单下架，托盘可增删", "待推广", "P3", "高", "备货员", "卡车出库主路径起点"),
    (6, "出库", "确认发车", "outbound-dispatch.html", "扣库存放行，发通知信", "待推广", "P3", "高", "发车/主管", "备货完成后发车"),
    (7, "出库", "自提出库", "self-pickup-outbound.html", "定位自提单→扫板放行→拍照；支持部分提货", "已上线使用", "P0", "—", "出库/门卫", "出库并行通道；唯一已在用"),
    (8, "出库", "快递出库", "express-outbound.html", "扫运单交接出库", "待推广", "P3", "中", "快递交接", "出库并行通道；可与卡车并行"),
    (9, "异常", "问题件处理", "problem.html", "多发少发纠错、未知货圈存", "待推广", "P4", "中高", "异常岗", "多与到仓环节并发；可随 P1 宣贯"),
    (10, "异常", "工单处理", "work-order.html", "执行网页端客服交办", "待推广", "P4", "中", "执行岗", "依赖 PC 工单联动"),
    (11, "异常", "整车退仓", "return.html", "已装车整批退库", "待推广", "P4", "中", "退仓岗", "发车后逆向；低频高风险"),
    (12, "异常", "单项退仓", "self-return.html", "定位 BOL 后登记退仓板", "待推广", "P4", "中", "退仓岗", "习惯接近自提"),
    (13, "异常", "板标重打", "supervisor.html", "污损板标原样重印", "待推广", "P4", "中", "主管授权", "全流程辅助；建议授权管控"),
    (14, "单据", "单据查询", "wms-pda.html 单据Tab", "按 REF 查货件详情", "待推广", "P1–P2", "中", "全员辅助", "全流程辅助；可随到仓培训"),
]
for i, row in enumerate(modules):
    r = 6 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=r, column=c, value=v)
    style_body(ws1, r, r, 10)
    if row[5] in status_fills:
        ws1.cell(row=r, column=6).fill = status_fills[row[5]]
    ph = str(row[6])[:2]
    if ph in phase_fills:
        ws1.cell(row=r, column=7).fill = phase_fills[ph]
    ws1.row_dimensions[r].height = 32

# --- B. P1 到仓操作要点 ---
r = 21
write_merged(ws1, r, 10, "B. P1 到仓确认 · 操作要点（培训提纲）", section_font)
r = 22
for i, h in enumerate(["步骤", "操作", "关键点", "常见错误", "", "", "", "", "", ""], 1):
    if h:
        ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 4)
ops = [
    ("1", "扫码/手输：柜号、提单号、系统单号、FBA、REF", "扫对单；留意重复到仓预警", "扫错单、重复确认"),
    ("2", "核对预估板数，填实际板数（必填）", "实收≠预估必须处理", "漏填、强行忽略不符"),
    ("3", "板数不符 → 登记问题件/快捷异常", "类型+描述+照片齐全", "只改数字不登记"),
    ("4", "拍到仓照片（≥1 张）", "拍清柜号/货况", "模糊、未拍柜号"),
    ("5", "可选备注 → 确认到仓", "确认前再核单号", "确认后难回溯"),
]
for i, row in enumerate(ops):
    rr = 23 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 4)
    ws1.row_dimensions[rr].height = 28

# --- C. P1 两周执行 ---
r = 29
write_merged(ws1, r, 10, "C. P1 到仓确认 · 两周执行计划", section_font)
r = 30
plan_h = ["天次", "动作", "负责人", "参与人", "产出/验收", "风险与应对", "完成标记", "", "", ""]
for i, h in enumerate(plan_h, 1):
    if h:
        ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 7)
plans = [
    ("D1", "定试点仓与种子用户(2–3人)；盘点 PDA 电量/扫码/拍照", "仓主管+IT", "种子用户", "名单、设备清单", "机不够→调自提闲时机", ""),
    ("D1–D2", "环境验收：单号可扫、照片可传、重复到仓有提示", "IT/实施", "产品", "验收记录", "码制不兼容→手输+贴码整改", ""),
    ("D2", "集中培训 30–45 分钟：全流程+板数不符走问题件", "培训讲师", "到仓岗", "签到；人均实操≥2单", "轮班→分两班培训", ""),
    ("D3–D5", "跟班带教，PC 仅兜底；每日收问题", "仓主管", "到仓岗", "每日问题清单", "高峰拥堵→加配 1 台", ""),
    ("D6–D7", "周复盘使用率/错单；出 SOP 一页纸", "产品+主管", "种子用户", "SOP v1、切换建议", "阻力大→延长试点", ""),
    ("W2 D8–10", "扩至全到仓岗；预告强制切换日", "仓主管", "全员", "PDA 占比目标≥80%", "抵触→一对一辅导", ""),
    ("W2 D11–12", "强制切换：到仓默认 PDA；PC 补录需审批", "仓主管", "全员", "切换公告执行", "故障→应急 PC/纸质", ""),
    ("W2 D13–14", "复盘并准备复制下一仓；启动 P2 打印机/库位准备", "产品+IT", "各仓主管", "复制包(SOP+清单)", "跨仓差异→按仓裁剪", ""),
]
for i, row in enumerate(plans):
    rr = 31 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 7)
    ws1.row_dimensions[rr].height = 32

# --- D. P1 KPI ---
r = 40
write_merged(ws1, r, 10, "D. P1 成功指标", section_font)
r = 41
for i, h in enumerate(["指标", "目标", "统计口径", "检视频率"], 1):
    ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 4)
for i, row in enumerate([
    ("到仓 PDA 占比", "W2 末≥80%，强制切换后≥95%", "PDA 到仓单 / 全部到仓单", "每日/每周"),
    ("到仓照片完整率", "100%（≥1 张）", "有照片单 / 到仓单", "每日"),
    ("板数不符登记率", "100%", "已登记问题件 / 实收≠预估", "每日"),
    ("重复到仓误操作", "趋近 0", "重复确认 / 到仓次数", "每周"),
    ("培训达标率", "到仓岗 100%", "考核通过 / 应训人数", "结束时"),
]):
    rr = 42 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 4)
    ws1.row_dimensions[rr].height = 24

# --- E. P2-P4 ---
r = 48
write_merged(ws1, r, 10, "E. P2–P4 后续阶段动作", section_font)
r = 49
for i, h in enumerate(["阶段", "模块", "启动前提", "关键动作", "成功标准", "建议周期", "完成标记"], 1):
    ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 7)
later = [
    ("P2", "拆柜打板", "P1 稳定；打印机到位", "联机培训；现场跟班；故障应急", "新板 100% 有打印记录", "1–1.5 周", ""),
    ("P2", "入库上架", "库位码覆盖关键库区", "扫板+库位培训；抽检准确率", "库位绑定率≥95%", "1–1.5 周", ""),
    ("P2", "私卡测量", "确认启用仓", "仅目标仓培训，衔接到上架", "启用仓先测后上架", "按需 3–5 天", ""),
    ("P3", "出库备货", "备货单可查；库存准", "拣货培训；错拣复盘", "错拣率下降", "1–1.5 周", ""),
    ("P3", "确认发车", "备货已 PDA 化；通知联调", "放行+二次确认培训", "扣库存及时；通知发出", "约 1 周", ""),
    ("P3", "快递出库", "运单扫码规则明确", "交接岗专项培训", "运单与出库匹配准确", "3–5 天", ""),
    ("P4", "问题件", "可与 P1 同步", "类型字典+照片规范+到仓联动", "当日登记≥90%", "3–5 天", ""),
    ("P4", "工单处理", "PC 工单可下发", "接收/执行/回传演练", "响应及时达标", "3–5 天", ""),
    ("P4", "整车/单项退仓", "权限与原因字典", "低频高风险演练+二次确认", "退仓可追溯、凭证齐", "约 1 周", ""),
    ("P4", "板标重打", "主管授权策略", "授权开通；审计抽查", "可追溯、无滥用", "2–3 天", ""),
]
for i, row in enumerate(later):
    rr = 50 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 7)
    if row[0] in phase_fills:
        ws1.cell(row=rr, column=1).fill = phase_fills[row[0]]
    ws1.row_dimensions[rr].height = 30

# --- F. 检查清单 ---
r = 61
write_merged(ws1, r, 10, "F. 上线前检查清单（每仓复制勾选）", section_font)
r = 62
for i, h in enumerate(["类别", "检查项", "标准", "结果(是/否)", "备注"], 1):
    ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 5)
checks = [
    ("账号", "账号开通且权限匹配岗位", "名单与岗位一致", "", ""),
    ("设备", "PDA 满足高峰并发；充电充足", "高峰排队不超 5 分钟", "", ""),
    ("扫码", "柜号/提单/板标/库位可识别", "抽测≥10 个样本成功", "", ""),
    ("拍照", "摄像头与上传正常", "试传 1 张成功", "", ""),
    ("打印", "打印机联机（P2 起）", "试打 1 张板标", "", ""),
    ("网络", "作业区网络稳定", "关键点可正常作业", "", ""),
    ("SOP", "一页纸操作卡已下发/张贴", "到仓岗可随时查看", "", ""),
    ("应急", "故障时 PC/纸质兜底已宣贯", "主管知晓启用条件", "", ""),
    ("反馈", "问题反馈群/表单已建", "响应人明确", "", ""),
    ("切换", "强制切换日已公告", "全员已知晓", "", ""),
]
for i, row in enumerate(checks):
    rr = 63 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 5)
    ws1.row_dimensions[rr].height = 24

# --- G. 培训课表简表 ---
r = 74
write_merged(ws1, r, 10, "G. 培训课表（简）", section_font)
r = 75
for i, h in enumerate(["场次", "对象", "时长", "内容", "考核"], 1):
    ws1.cell(row=r, column=i, value=h)
style_header(ws1, r, 5)
for i, row in enumerate([
    ("启动会", "主管+种子用户", "30 分钟", "目标、节奏、切换日、反馈通道", "确认名单排期"),
    ("到仓实操", "到仓岗", "45 分钟", "扫码→板数→照片→确认；不符走问题件", "实操 2 单通过"),
    ("入库闭环", "拆柜/上架岗", "60 分钟", "打板打印+上架绑库位", "打板+上架各 1 次"),
    ("出库扩展", "备货/发车/快递", "60 分钟", "备货、发车、快递扫运单", "模拟单走通"),
    ("异常逆向", "异常/退仓/主管", "45 分钟", "问题件、工单、退仓、重打权限", "案例演练"),
]):
    rr = 76 + i
    for c, v in enumerate(row, 1):
        ws1.cell(row=rr, column=c, value=v)
    style_body(ws1, rr, rr, 5)
    ws1.row_dimensions[rr].height = 26

set_widths(ws1, [10, 14, 36, 28, 36, 12, 10, 14, 12, 22])
ws1.freeze_panes = "A6"

out = os.path.join(os.path.dirname(__file__), "PDA功能推广使用计划.xlsx")
wb.save(out)
print("OK", out)
